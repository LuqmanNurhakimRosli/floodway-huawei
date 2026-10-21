import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.oxml.ns import qn
from lxml import etree

TPL = '/root/.claude/uploads/6ba3729a-45ea-58fa-8efd-cad6f3b8411e/1514225d-SDP_University_Hackathon_TOWNSHIP-OF-TOMORROW-Template.pptx'
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'FloodWay_SDP_Pitch_Deck.pptx')

RED, NAVY, TEXT, LIGHT, MUTED = 'C8102E', '152D47', '2B3A4A', 'EEF1F5', '566475'
WHITE, TINT, BLUE, SOFTBLUE, YEL = 'FFFFFF', 'FBE9EC', '1E7FD0', 'CFE3F6', 'FFF1A8'
FONT = 'Arial'


def rgb(h):
    return RGBColor.from_string(h)


prs = Presentation(TPL)
# drop all template slides (keep layouts/master)
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)

LAY_SECTION = prs.slide_layouts.get_by_name('Section Header')
LAY_TITLE = prs.slide_layouts.get_by_name('3_Title Slide')


# ---------- helpers ----------
def set_runs(p, runs, size, color, bold, font=FONT, italic=False):
    for r in runs:
        if isinstance(r, str):
            r = (r, {})
        t, o = r
        run = p.add_run()
        run.text = t
        f = run.font
        f.name = font
        f.size = Pt(o.get('size', size))
        f.bold = o.get('bold', bold)
        f.italic = o.get('italic', italic)
        f.color.rgb = rgb(o.get('color', color))
        if o.get('link'):
            run.hyperlink.address = o['link']
        if o.get('hl'):
            rPr = run._r.get_or_add_rPr()
            h = etree.SubElement(rPr, qn('a:highlight'))
            c = etree.SubElement(h, qn('a:srgbClr'))
            c.set('val', o['hl'])
            # highlight must precede latin; move to correct spot
            rPr.remove(h)
            rPr.insert(1 if rPr.find(qn('a:solidFill')) is not None else 0, h)


def tb(slide, x, y, w, h, paras, size=14, color=TEXT, bold=False, align='l', anchor='t',
       name=None, font=FONT, space_after=0, line=None, margin=0.0, italic=False):
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str):
        paras = [paras]
    first = True
    for item in paras:
        if isinstance(item, (str, tuple)):
            item = {'runs': [item]}
        elif isinstance(item, list):
            item = {'runs': item}
        elif 'text' in item:
            item = dict(item)
            item['runs'] = [item.pop('text')]
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[item.get('align', align)]
        p.space_after = Pt(item.get('space_after', space_after))
        if item.get('space_before'):
            p.space_before = Pt(item['space_before'])
        if line or item.get('line'):
            p.line_spacing = item.get('line', line)
        set_runs(p, item['runs'], item.get('size', size), item.get('color', color),
                 item.get('bold', bold), item.get('font', font), item.get('italic', italic))
        if item.get('bullet'):
            pPr = p._p.get_or_add_pPr()
            pPr.set('marL', '171450')
            pPr.set('indent', '-171450')
            bf = etree.SubElement(pPr, qn('a:buFont'))
            bf.set('typeface', 'Arial')
            bc = etree.SubElement(pPr, qn('a:buChar'))
            bc.set('char', '•')
    return s


def box(slide, x, y, w, h, fill=LIGHT, line=None, radius=0.05, shape=None, name=None, alpha=None, lw=1.0):
    st = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    s = slide.shapes.add_shape(st, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    if radius and st == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            clr = s.fill._xPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
            a = etree.SubElement(clr, qn('a:alpha'))
            a.set('val', str(int(alpha * 1000)))
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    s.text_frame.text = ''
    return s


def badge(slide, x, y, n, d=0.38, fill=RED, size=14):
    s = box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL, radius=0, name=f'Badge {n}')
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    set_runs(p, [str(n)], size, WHITE, True)
    return s


def line(slide, x1, y1, x2, y2, color=NAVY, w=3):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    return c


def add_title_ph(slide, text, x=0.7, y=0.62, w=10.2, h=0.62, size=26, color=NAVY):
    """Real title placeholder (inherits from master) so outline view/screen readers see titles."""
    spTree = slide.shapes._spTree
    sid = max([int(e.get('id')) for e in spTree.iter() if e.tag.endswith('cNvPr')] + [1]) + 1
    xml = f'''<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
      xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
      <p:nvSpPr><p:cNvPr id="{sid}" name="Title {sid}"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>
      <p:nvPr><p:ph type="title"/></p:nvPr></p:nvSpPr>
      <p:spPr><a:xfrm><a:off x="{int(Inches(x))}" y="{int(Inches(y))}"/><a:ext cx="{int(Inches(w))}" cy="{int(Inches(h))}"/></a:xfrm></p:spPr>
      <p:txBody><a:bodyPr lIns="0" tIns="0" rIns="0" bIns="0" anchor="ctr"><a:normAutofit/></a:bodyPr><a:lstStyle/>
      <a:p><a:pPr algn="l"/><a:r><a:rPr lang="en-GB" sz="{size*100}" b="1" dirty="0"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill>
      <a:latin typeface="{FONT}"/><a:cs typeface="{FONT}"/></a:rPr><a:t>{text}</a:t></a:r></a:p></p:txBody></p:sp>'''
    spTree.append(etree.fromstring(xml))



def set_notes(slide, text):
    ns = slide.notes_slide
    tf = ns.notes_text_frame
    if tf is None:
        spTree = ns.shapes._spTree
        xml = ('<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
               'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:nvSpPr><p:cNvPr id="3" name="Notes Placeholder 2"/>'
               '<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr><p:ph type="body" idx="1"/></p:nvPr></p:nvSpPr>'
               '<p:spPr><a:xfrm><a:off x="685800" y="4400550"/><a:ext cx="5486400" cy="3600450"/></a:xfrm></p:spPr>'
               '<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:endParaRPr lang="en-GB"/></a:p></p:txBody></p:sp>')
        spTree.append(etree.fromstring(xml))
        tf = ns.notes_text_frame
    tf.text = text

page_no = [0]


def content_slide(label, title, source=None, notes=None, appendix=False):
    s = prs.slides.add_slide(LAY_SECTION)
    page_no[0] += 1
    tb(s, 0.7, 0.3, 9.5, 0.28, [label], size=11, color=RED, bold=True, name='Section label')
    add_title_ph(s, title)
    if source:
        tb(s, 0.7, 6.98, 11.2, 0.3, [source], size=10, color=MUTED, name='Source note', anchor='m')
    tb(s, 12.2, 6.98, 0.45, 0.3, [str(page_no[0]) if not appendix else 'A' + str(page_no[0] - 10)],
       size=10, color=MUTED, align='r', name='Page number', anchor='m')
    if notes:
        set_notes(s, notes)
    return s


def fill_tag(text):
    return (text, {'hl': YEL, 'color': NAVY, 'bold': True})


# =====================================================================
# 1. TITLE
# =====================================================================
s = prs.slides.add_slide(LAY_TITLE)
page_no[0] += 1
for ph in list(s.placeholders):
    ph._element.getparent().remove(ph._element)
add_title_ph(s, 'FloodWay', x=1.4, y=2.15, w=10.5, h=1.0, size=60, color=WHITE)
s.shapes[-1].text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
tb(s, 1.4, 3.3, 10.5, 0.6, ['Predict. Verify. Evacuate safely.'], size=28, color=WHITE, bold=True, align='c', name='Tagline')
tb(s, 1.4, 4.2, 10.5, 1.2, [
    {'text': 'AI flood preparedness and evacuation for Malaysian townships', 'size': 18, 'space_after': 8},
    {'text': 'Sime Darby Property University Hackathon 2026  ·  Township of Tomorrow  ·  Planet track', 'size': 14},
], color=WHITE, align='c', name='Subtitle')
tb(s, 1.4, 6.35, 10.5, 0.4, ['Live prototype: floodways.netlify.app'], size=14, color=WHITE, align='c', name='Prototype link')
set_notes(s, (
    'Open with the one-line promise: FloodWay turns a flood warning into a safe decision for one street. '
    'Track choice: Planet (climate resilience), with a life-safety angle that also fits People.'))

# =====================================================================
# 2. TEAM
# =====================================================================
s = content_slide('TEAM INFORMATION', 'Team information',
                  source='Fill the highlighted fields before exporting the PDF. Names must match the registration.',
                  notes='Introduce the team and say who did what. Mention that FloodWay has been developed and improved since March 2025.')
rows = [('Team name', fill_tag('[Team name]')), ('Team leader', fill_tag('[Full name]')),
        ('University', 'Universiti Teknologi MARA (UiTM), Shah Alam  ·  Faculty of Computer and Mathematical Sciences'),
        ('Team members (2–5)', fill_tag('[Name 1]  ·  [Name 2]  ·  [Name 3]  ·  [Name 4]')),
        ('Email', fill_tag('[Leader email]')), ('Contact number', fill_tag('[Leader phone]'))]
y = 1.55
for k, v in rows:
    box(s, 0.7, y, 6.6, 0.68, fill=LIGHT, name=f'Row {k}')
    tb(s, 0.9, y, 2.1, 0.68, [k], size=13, bold=True, color=NAVY, anchor='m')
    tb(s, 3.0, y, 4.2, 0.68, [[v]], size=13, anchor='m')
    y += 0.78
tb(s, 7.7, 1.55, 4.9, 0.4, ['What the team brings'], size=18, bold=True, color=NAVY)
items = [('Built and running', 'A live, installable PWA with prediction, photo verification and risk-aware routing.'),
         ('Improved since March 2025', 'FloodWay began as a UiTM competition project and has been rebuilt and extended since (project report v2.2, Feb 2026).'),
         ('Right skills', 'Machine-learning deployment, geospatial analysis and full-stack web development.')]
y = 2.1
for i, (h, t) in enumerate(items, 1):
    badge(s, 7.7, y + 0.03, i)
    tb(s, 8.25, y, 4.35, 1.25, [{'text': h, 'bold': True, 'color': NAVY, 'size': 14, 'space_after': 2},
                                {'text': t, 'size': 12}], size=12)
    y += 1.4

# =====================================================================
# 3. PROBLEM + TARGET USER
# =====================================================================
s = content_slide('1 PROBLEM  ·  2 TARGET USER', 'Residents cannot answer four questions in a flood',
                  source='Sources: DOSM special reports (2021, 2024, 2025); The Vibes (police toll); Sinar Harian (May 2024); SDP Sustainability Report 2023. Full list: Appendix A3.',
                  notes='Problem: floods recur, losses are large, and at street level people still need four answers. '
                        'Be careful with the Bukit Raja line: SDP already works on flood mitigation; we offer a digital layer.')
stats = [('RM6.1bn', 'flood losses, Dec 2021–Jan 2022 floods (DOSM)'),
         ('54', 'deaths in the Dec 2021 floods (police count, 5 Jan 2022)'),
         ('RM0.6–0.9bn', 'flood losses a year, 2023–2025 (DOSM)')]
x = 0.7
for big, lab in stats:
    box(s, x, 1.5, 2.3, 1.5, fill=LIGHT)
    tb(s, x + 0.15, 1.6, 2.05, 0.6, [big], size=21, bold=True, color=RED, anchor='m')
    tb(s, x + 0.15, 2.2, 2.0, 0.75, [lab], size=11, color=TEXT)
    x += 2.45
tb(s, 0.7, 3.25, 7.2, 0.35, ['Four questions residents cannot answer today'], size=16, bold=True, color=NAVY)
qs = ['When will floodwater reach my street?', 'Which road is still passable?',
      'Is this flood photo real, or an old one?', 'Which relief centre (PPS) is open and reachable?']
y = 3.7
for i, q in enumerate(qs, 1):
    badge(s, 0.7, y, i, d=0.36, size=13)
    tb(s, 1.2, y, 6.7, 0.36, [q], size=14, anchor='m')
    y += 0.5
box(s, 0.7, 5.8, 7.2, 1.05, fill=TINT)
tb(s, 0.9, 5.85, 6.8, 0.95, [[('Why townships: ', {'bold': True, 'color': RED}),
                              'Bandar Bukit Raja residents reported repeated flash floods since Dec 2021, and SDP lists flood exposure as a key risk. '
                              'Police and NADMA also keep warning that old flood photos are recirculated.']],
   size=12, anchor='m')
tb(s, 8.3, 1.5, 3.0, 0.4, ['Target users'], size=18, bold=True, color=NAVY)
users = [('Residents and commuters', 'Families, older residents and drivers in flood-prone townships who must decide when to leave and which way to go.'),
         ('Township management and committees', 'Need verified, street-level information to brief residents and run drills.'),
         ('Responders and authorities', 'NADMA, Bomba, APM and local councils need verified reports and shelter status.')]
y = 2.0
for h, t in users:
    box(s, 8.3, y, 4.3, 1.5, fill=LIGHT)
    tb(s, 8.5, y + 0.1, 3.9, 1.35, [{'text': h, 'bold': True, 'color': NAVY, 'size': 14, 'space_after': 3},
                                    {'text': t, 'size': 12}])
    y += 1.65

# =====================================================================
# 4. SOLUTION + VISUAL CONCEPT
# =====================================================================
s = content_slide('3 SOLUTION  ·  VISUAL CONCEPT', 'One flow: predict risk, verify reports, route to safety',
                  source='Illustrative mockups of the app flow; example values use the prototype\'s output format. Decision-support only: official warnings remain with JPS and MetMalaysia.',
                  notes='Walk the three phones left to right. Demo video shows the real app.')
steps = [('Predict', 'A flood probability and risk class for your location, from recent rainfall, forecast and flood history.'),
         ('Verify', 'Resident photo reports are checked by AI. Low-confidence reports go to a human moderator before they steer anyone.'),
         ('Route', 'Evacuation routes avoid verified flooded roads and end at the nearest reachable relief centre (PPS).')]
y = 1.55
for i, (h, t) in enumerate(steps, 1):
    badge(s, 0.7, y + 0.02, i)
    tb(s, 1.25, y, 3.7, 1.35, [{'text': h, 'bold': True, 'color': NAVY, 'size': 17, 'space_after': 2},
                               {'text': t, 'size': 13}])
    y += 1.5
box(s, 0.7, 6.1, 4.3, 0.65, fill=LIGHT)
tb(s, 0.85, 6.1, 4.0, 0.65, [[('Try it: ', {'bold': True, 'color': RED}), ('floodways.netlify.app', {'link': 'https://floodways.netlify.app/', 'color': BLUE})],
                              {'text': 'Installable PWA, no app-store download', 'size': 11}], size=13, anchor='m')

PX, PY, PW, PH = 5.55, 1.7, 2.15, 4.35
caps = ['1  Predict', '2  Verify', '3  Route']
for k in range(3):
    px = PX + k * (PW + 0.3)
    box(s, px, PY, PW, PH, fill=NAVY, radius=0.1, name=f'Phone {k+1}')
    box(s, px + 0.1, PY + 0.12, PW - 0.2, PH - 0.24, fill=WHITE, radius=0.06, name=f'Screen {k+1}')
    tb(s, px, PY + PH + 0.1, PW, 0.35, [caps[k]], size=14, bold=True, color=NAVY, align='c')
sx = lambda k: PX + k * (PW + 0.3) + 0.22
sw = PW - 0.44
# phone 1
k = 0
tb(s, sx(k), PY + 0.3, sw, 0.3, ['Your street · Shah Alam'], size=10, bold=True, color=NAVY)
p = box(s, sx(k), PY + 0.7, sw, 0.5, fill=RED, radius=0.3)
tb(s, sx(k), PY + 0.7, sw, 0.5, ['HIGH RISK'], size=15, bold=True, color=WHITE, align='c', anchor='m')
tb(s, sx(k), PY + 1.35, sw, 0.6, [{'text': 'Flood probability', 'size': 10, 'color': MUTED}, {'text': '82%', 'size': 26, 'bold': True, 'color': NAVY}])
tb(s, sx(k), PY + 2.1, sw, 0.5, [{'text': 'Estimated time to flood', 'size': 10, 'color': MUTED}, {'text': '4–6 hours', 'size': 14, 'bold': True, 'color': NAVY}])
tb(s, sx(k), PY + 2.8, sw, 0.25, ['Rainfall'], size=10, color=MUTED)
for j, (lab, hgt) in enumerate([('24h', 0.35), ('48h', 0.7), ('72h', 0.95)]):
    bx = sx(k) + 0.1 + j * 0.55
    box(s, bx, PY + 4.05 - hgt, 0.38, hgt, fill=BLUE, radius=0)
    tb(s, bx - 0.05, PY + 3.62 + 0.08 + 0.35 - 0.0, 0.48, 0.2, [lab], size=9, color=MUTED, align='c') if False else None
    tb(s, bx - 0.05, PY + 4.05, 0.48, 0.2, [lab], size=9, color=MUTED, align='c')
# phone 2
k = 1
tb(s, sx(k), PY + 0.3, sw, 0.3, ['Report a flood'], size=10, bold=True, color=NAVY)
box(s, sx(k), PY + 0.7, sw, 1.5, fill='B8C2CE', radius=0.05, name='Photo placeholder')
box(s, sx(k), PY + 1.5, sw, 0.7, fill=BLUE, radius=0, alpha=60, name='Water in photo')
tb(s, sx(k), PY + 0.75, sw, 0.3, ['[ user photo ]'], size=10, color=NAVY, align='c')
chips = [('Water detected', NAVY), ('Confidence 92', NAVY), ('HIGH · 0.5–1.0 m deep', RED)]
yy = PY + 2.38
for t, c in chips:
    box(s, sx(k), yy, sw, 0.32, fill=LIGHT, radius=0.4)
    tb(s, sx(k) + 0.08, yy, sw - 0.1, 0.32, [t], size=10, bold=True, color=c, anchor='m')
    yy += 0.38
box(s, sx(k), yy + 0.05, sw, 0.42, fill=NAVY, radius=0.3)
tb(s, sx(k), yy + 0.05, sw, 0.42, ['Auto-verified, mapped'], size=10, bold=True, color=WHITE, align='c', anchor='m')
tb(s, sx(k), yy + 0.5, sw, 0.25, ['Else: human review'], size=9, color=MUTED, align='c')
# phone 3
k = 2
tb(s, sx(k), PY + 0.3, sw, 0.3, ['Safe route'], size=10, bold=True, color=NAVY)
box(s, sx(k), PY + 0.7, sw, 2.5, fill='E4EAF0', radius=0.04, name='Map')
for (a, b, c, d) in [(0.05, 0.4, 1.7, 0.4), (0.05, 1.3, 1.7, 1.3), (0.05, 2.15, 1.7, 2.15), (0.35, 0.1, 0.35, 2.4), (1.0, 0.1, 1.0, 2.4), (1.4, 0.1, 1.4, 2.4)]:
    line(s, sx(k) + a, PY + 0.7 + b, sx(k) + c, PY + 0.7 + d, color='FFFFFF', w=5)
box(s, sx(k) + 0.55, PY + 0.7 + 0.65, 0.85, 0.95, fill=BLUE, shape=MSO_SHAPE.OVAL, radius=0, alpha=55, name='Verified flood zone')
pts = [(0.35, 2.15), (0.35, 1.3), (0.35, 0.4)]
line(s, sx(k) + 0.35, PY + 0.7 + 2.15, sx(k) + 0.35, PY + 0.7 + 0.4, color=RED, w=4)
line(s, sx(k) + 0.35, PY + 0.7 + 0.4, sx(k) + 1.4, PY + 0.7 + 0.4, color=RED, w=4)
box(s, sx(k) + 0.27, PY + 0.7 + 2.07, 0.16, 0.16, fill=NAVY, shape=MSO_SHAPE.OVAL, radius=0)
box(s, sx(k) + 1.2, PY + 0.7 + 0.18, 0.46, 0.4, fill=RED, radius=0.15)
tb(s, sx(k) + 1.2, PY + 0.7 + 0.18, 0.46, 0.4, ['PPS'], size=10, bold=True, color=WHITE, align='c', anchor='m')
tb(s, sx(k) + 0.55, PY + 0.7 + 0.95, 0.85, 0.3, ['flooded'], size=9, bold=True, color=NAVY, align='c')
tb(s, sx(k), PY + 3.35, sw, 0.7, [{'text': 'Avoids verified flooded roads', 'size': 11, 'bold': True, 'color': NAVY},
                                  {'text': 'to the nearest reachable relief centre', 'size': 10, 'color': MUTED}])

# =====================================================================
# 5. INNOVATION
# =====================================================================
s = content_slide('4 INNOVATION', 'Others cover parts of the chain; we add the last mile',
                  source='Checked Oct 2026 against official help pages and news. "Not found" means searched, not proof of absence. Evidence links: Appendix A3.',
                  notes='Be generous to competitors: Waze and Google do take flood reports and closures; MyDIMS now routes to shelters. '
                        'Our claim is narrow: street-level timing, AI-checked photos, and flood-aware routing to a PPS. We say what is proven and what is a target.')
hdr = ['', 'Waze', 'Google Maps and Flood Hub', 'Official Malaysian apps (MyPublicInfoBanjir, MyDIMS)', 'FloodWay']
data = [
    ['Flood forecast, 1+ day ahead', 'Not found', 'Yes · river basins, up to 7 days', 'Yes · JPS bulletins, district and basin level', 'Prototype · district-level model, being validated'],
    ['Street-level flood timing', 'Not found', 'Not found · flash model is 20 km, 24 h', 'Not found', 'Target · the pilot will test it'],
    ['Resident flood reports', 'Yes · hazard and closure reports', 'Partial · "flooded road" test, Malaysia unconfirmed', 'Yes · MyDIMS, verified before broadcast', 'Yes · geotagged photo reports'],
    ['AI check of photo reports', 'Partial · crowd votes only', 'Not found', 'Method not described', 'Yes · Gemini vision; accuracy being measured'],
    ['Route to an open relief centre (PPS)', 'Partial · staff-run crisis pins', 'Not found', 'Yes · MyDIMS navigation', 'Yes · nearest reachable PPS'],
    ['Flood-aware route computation', 'Partial · closures only', 'Not found', 'Method not described', 'Yes · simulation only; field test in pilot'],
]
colw = [2.55, 1.95, 2.35, 2.65, 2.4]
tbl_shape = s.shapes.add_table(len(data) + 1, 5, Inches(0.7), Inches(1.75), Inches(sum(colw)), Inches(3.95))
tbl_shape.name = 'Comparison matrix'
tbl = tbl_shape.table
# strip default style banding
tblPr = tbl._tbl.tblPr
tblPr.set('firstRow', '0')
tblPr.set('bandRow', '0')
for i, w in enumerate(colw):
    tbl.columns[i].width = Inches(w)
tbl.rows[0].height = Inches(0.6)
for r in range(1, len(data) + 1):
    tbl.rows[r].height = Inches(0.56)


def cell(c, runs, fill, color=TEXT, size=11, bold=False, align='l'):
    c.fill.solid()
    c.fill.fore_color.rgb = rgb(fill)
    c.margin_left = c.margin_right = Inches(0.08)
    c.margin_top = c.margin_bottom = Inches(0.04)
    c.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = c.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER}[align]
    set_runs(p, runs, size, color, bold)


for j, h in enumerate(hdr):
    cell(tbl.cell(0, j), [h], RED if j == 4 else NAVY, WHITE, 11, True)
for i, row in enumerate(data, 1):
    for j, v in enumerate(row):
        c = tbl.cell(i, j)
        if j == 0:
            cell(c, [v], WHITE, NAVY, 11, True)
            continue
        head, _, rest = v.partition(' · ')
        if j == 4:
            cell(c, [(head, {'bold': True, 'color': RED})] + ([(' · ' + rest, {})] if rest else []), TINT, NAVY, 11)
        elif v.startswith('Yes'):
            cell(c, [(head, {'bold': True})] + ([(' · ' + rest, {})] if rest else []), 'D6E4F2', NAVY, 11)
        elif v.startswith('Partial'):
            cell(c, [(head, {'bold': True})] + ([(' · ' + rest, {})] if rest else []), LIGHT, NAVY, 11)
        else:
            cell(c, [v], WHITE, MUTED, 11)
# thin cell borders
for r in tbl.rows:
    for c in r.cells:
        tcPr = c._tc.get_or_add_tcPr()
        for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
            ln = etree.SubElement(tcPr, qn(tag))
            ln.set('w', '9525')
            sf = etree.SubElement(ln, qn('a:solidFill'))
            cc = etree.SubElement(sf, qn('a:srgbClr'))
            cc.set('val', 'C9D1DA')
        # borders must come before fill in schema order: move fill to end
        fills = [e for e in tcPr if e.tag in (qn('a:solidFill'),)]
        for f in fills:
            tcPr.remove(f)
            tcPr.append(f)
yb = 5.9
tb(s, 0.7, yb, 11.9, 0.3, ['Where FloodWay is different, if the pilot proves it'], size=14, bold=True, color=NAVY)
diffs = ['Finer than a 20 km forecast cell', 'Photo reports checked before they steer routes', 'Flood-aware routing that ends at a relief centre']
x = 0.7
for d in diffs:
    box(s, x, yb + 0.38, 3.85, 0.5, fill=LIGHT)
    tb(s, x + 0.12, yb + 0.38, 3.65, 0.5, [d], size=12, bold=True, color=NAVY, anchor='m')
    x += 4.02

# =====================================================================
# 6. HOW IT WORKS
# =====================================================================
s = content_slide('5 HOW IT WORKS', 'How it works: three engines on disclosed public data',
                  source='Stack: React, TypeScript, Leaflet and OpenStreetMap on the front end; FastAPI, scikit-learn, Firebase and Netlify on the back end. Output examples are illustrative.',
                  notes='Explain each engine in one breath. Stress that all training data is public and listed in Appendix A1; no SDP data was used.')
cards = [
    ('1', 'Prediction engine', [
        ('In', 'Rainfall over the past 24, 48 and 72 hours, location, flood history of the area, season'),
        ('How', 'Random Forest classifier (scikit-learn), trained on public rainfall and flood-incident data'),
        ('Out', 'Flood probability and a risk class (Low, Moderate, High, Critical) with a time-to-flood band')],
     '{"floodProbability": 0.82,\n "riskLevel": "HIGH",\n "estimatedTimeToFlood": "4-6 hours"}'),
    ('2', 'Truth filter', [
        ('In', 'A geotagged photo from a resident'),
        ('How', 'Gemini Flash vision model scores water, severity and confidence. Above the threshold: auto-mapped. Below: human moderator queue (planned)'),
        ('Out', 'A verified flood pin on the map, or a held report')],
     '{"waterDetected": true,\n "confidence": 92,\n "severity": "HIGH"}'),
    ('3', 'Risk-aware routing', [
        ('In', 'OpenStreetMap roads, verified flood pins, relief centre (PPS) locations'),
        ('How', 'Road graph with risk penalties; verified flooded roads are blocked; OSRM computes the route'),
        ('Out', 'Turn-by-turn route to the nearest reachable PPS; a "no safe route" fallback is planned')],
     'edge weight =\n  distance + risk penalty'),
]
x = 0.7
for n, title, rows_, code in cards:
    box(s, x, 1.75, 3.85, 3.55, fill=LIGHT)
    badge(s, x + 0.15, 1.87, n)
    tb(s, x + 0.65, 1.87, 3.1, 0.38, [title], size=16, bold=True, color=NAVY, anchor='m')
    paras = []
    for lab, t in rows_:
        paras.append({'runs': [(lab.upper() + '  ', {'bold': True, 'color': RED, 'size': 10}), (t, {})], 'size': 11.5, 'space_after': 5})
    tb(s, x + 0.2, 2.4, 3.45, 2.0, paras, size=11.5)
    box(s, x + 0.2, 4.45, 3.45, 0.7, fill=NAVY, radius=0.06)
    tb(s, x + 0.3, 4.45, 3.3, 0.7, [code], size=10, color=WHITE, font='Courier New', anchor='m')
    x += 4.025
box(s, 0.7, 5.45, 11.9, 1.4, fill=WHITE, line='C9D1DA')
tb(s, 0.9, 5.52, 11.5, 0.3, ['Training data: public sources only, full register in Appendix A1'], size=13, bold=True, color=NAVY)
ds = [('Dataset 1', 'Malaysian rainfall 2000–2010: state, district and monthly rainfall patterns'),
      ('Dataset 2', 'Malaysian flood incidents 2015–2021 from JPS annual reports: locations, dates, geocoded coordinates, used as supervised labels'),
      ('Dataset 3', 'Real-time rainfall forecast: integration ready, used at prediction time')]
x = 0.9
for h, t in ds:
    tb(s, x, 5.9, 3.7, 0.9, [{'runs': [(h + '  ', {'bold': True, 'color': RED}), t]}], size=11.5)
    x += 3.85

# =====================================================================
# 7. EVIDENCE
# =====================================================================
s = content_slide('5 HOW IT WORKS  ·  EVIDENCE', 'Proof so far: working prototype, honest baseline',
                  source='Model accuracy from the team\'s initial hold-out split (prototype). Survey is a convenience sample. Method and limits: Appendix A2.',
                  notes='Do not oversell 74.55%. Say what it is: an initial split. The evaluation plan is in Appendix A2. '
                        'If asked about the Random Forest versus Decision Tree gap: it is small and we are re-testing.')
cd = CategoryChartData()
cd.categories = ['Random Forest', 'Decision Tree', 'KNN', 'Ensemble', 'Logistic Regression']
cd.add_series('Accuracy (%)', (74.55, 72.72, 66.06, 65.45, 53.93))
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.7), Inches(1.9), Inches(6.5), Inches(3.55), cd)
gf.name = 'Model accuracy chart'
ch = gf.chart
ch.has_legend = False
ch.has_title = False
ch.font.size = Pt(12)
ch.font.name = FONT
ch.font.color.rgb = rgb(TEXT)
plot = ch.plots[0]
plot.gap_width = 55
plot.has_data_labels = True
dl = plot.data_labels
dl.number_format = '0.00"%"'
dl.number_format_is_linked = False
dl.position = XL_LABEL_POSITION.OUTSIDE_END
dl.font.size = Pt(12)
dl.font.bold = True
dl.font.color.rgb = rgb(NAVY)
va = ch.value_axis
va.minimum_scale, va.maximum_scale = 0, 100
va.visible = False
va.has_major_gridlines = False
ca = ch.category_axis
ca.reverse_order = True
ca.tick_labels.font.size = Pt(12)
ca.format.line.color.rgb = rgb('C9D1DA')
ser = plot.series[0]
for i in range(5):
    pt = ser.points[i]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = rgb(RED if i == 0 else '9AA8B8')
tb(s, 0.7, 1.5, 6.5, 0.35, ['Model comparison: accuracy on an initial hold-out split'], size=15, bold=True, color=NAVY)
tb(s, 0.7, 5.5, 6.5, 1.35, [
    {'runs': [('Random Forest is selected, but the gap to the Decision Tree is under two points. ', {'bold': True, 'color': NAVY}),
              'We are re-testing on a time-blocked split against a no-flood and a climatology baseline. Flood share of the test set: ', fill_tag('[TBD %]'), '.']}],
   size=12)
cards = [
    ('Survey: 100 respondents', ['ASEAN youth and residents, convenience sample (not representative). Respondents rated AI early warning a priority: ', fill_tag('[insert: __ of 100 ...]')]),
    ('Working prototype', ['Live PWA with prediction, photo verification and risk-aware routing. Routing tested in synthetic scenarios: simulation, not field-tested.']),
    ('Measured next, before the pilot', ['Time-blocked split, PR-AUC and Brier score against baselines, photo-verification false-accept rate, blocked-road routing test.']),
]
y = 1.75
hh = [1.55, 1.4, 1.4]
for (h, t), hgt in zip(cards, hh):
    box(s, 7.6, y, 5.0, hgt, fill=LIGHT)
    tb(s, 7.8, y + 0.1, 4.6, hgt - 0.15, [{'text': h, 'bold': True, 'color': NAVY, 'size': 14, 'space_after': 3},
                                         {'runs': t, 'size': 12}])
    y += hgt + 0.12

# =====================================================================
# 8. IMPACT
# =====================================================================
s = content_slide('6 IMPACT', 'Impact: earlier warnings help, and we will measure ours',
                  source='Sources: Global Commission on Adaptation 2019 via WMO; UNDRR/WMO 2022; DOSM special reports 2023–2025. KPIs are targets, not results. Appendix A3.',
                  notes='We use only verified headline numbers. We do not claim a percentage improvement for FloodWay itself; the pilot measures it.')
tb(s, 0.7, 1.5, 4.4, 0.35, ['Why early warning matters (verified)'], size=15, bold=True, color=NAVY)
facts = [('~30%', 'less damage when warning comes 24 hours ahead (Global Commission on Adaptation 2019, via WMO)'),
         ('8×', 'higher disaster mortality in countries with limited early warning (UNDRR/WMO 2022)'),
         ('RM0.6–0.9bn', 'flood losses a year in 2023–2025 (DOSM)')]
y = 1.95
for big, lab in facts:
    box(s, 0.7, y, 4.4, 1.38, fill=LIGHT)
    tb(s, 0.9, y + 0.08, 4.0, 0.55, [big], size=26, bold=True, color=RED, anchor='m')
    tb(s, 0.9, y + 0.65, 4.0, 0.7, [lab], size=11.5)
    y += 1.52
tb(s, 5.5, 1.5, 7.1, 0.35, ['What we will measure in a township pilot (targets)'], size=15, bold=True, color=NAVY)
kp = [('Alert lead time', 'Minutes between a FloodWay alert and the first reported road flooding'),
      ('Report quality', 'False-accept and false-reject rate against a labelled photo set'),
      ('Route safety', 'Share of test routes that avoid verified flooded roads'),
      ('Reach', 'Registered residents and drills completed in the pilot township'),
      ('Time to safety', 'Time to reach a PPS in a drill, compared with a map-app route')]
y = 1.95
for i, (k_, t) in enumerate(kp):
    box(s, 5.5, y, 7.1, 0.56, fill=LIGHT if i % 2 == 0 else WHITE, radius=0.04)
    tb(s, 5.65, y, 1.85, 0.56, [k_], size=12, bold=True, color=NAVY, anchor='m')
    tb(s, 7.5, y, 5.0, 0.56, [t], size=12, anchor='m')
    y += 0.6
box(s, 5.5, 5.05, 7.1, 1.32, fill=TINT)
tb(s, 5.7, 5.12, 6.7, 1.2, [
    {'text': 'Value for SDP townships (assumptions to test in the pilot)', 'bold': True, 'color': RED, 'size': 13, 'space_after': 4},
    {'text': 'A reusable digital layer on top of wetlands and ponds', 'bullet': True, 'space_after': 2},
    {'text': 'Clearer resident communication and fewer unsafe trips', 'bullet': True, 'space_after': 2},
    {'text': 'One playbook that can repeat across townships. Supports SDG 11 and SDG 13', 'bullet': True},
], size=12)

# =====================================================================
# 9. FEASIBILITY + IMPLEMENTATION
# =====================================================================
s = content_slide('7 FEASIBILITY  ·  IMPLEMENTATION', 'Feasibility: a pilot-ready PWA and a one-township plan',
                  source='Pilot township is a proposal to be agreed with SDP. Wetland figure: SDP Sustainability Report 2023. Privacy: PDPA 2010 as amended 2024.',
                  notes='The pilot is shadow mode first: alerts reviewed by humans before residents see them. '
                        'We ask SDP for township contacts and access to existing sensors or camera feeds, and do not assume any SDP data.')
box(s, 0.7, 1.75, 11.9, 0.8, fill=TINT)
tb(s, 0.9, 1.75, 11.5, 0.8, [[('Proposed pilot: Bandar Bukit Raja, Klang (subject to SDP agreement). ', {'bold': True, 'color': RED}),
                              'FloodWay adds a digital layer for residents on top of SDP\'s physical flood mitigation, including its 180 acres of wetlands.']],
   size=13, anchor='m')
phases = [('Weeks 0–4', 'Fix and measure', 'Period-matched rainfall retrain, labelled photo test set, blocked-road routing test'),
          ('Weeks 5–8', 'Connect', 'Township contacts, PPS list, JPS and MetMalaysia data access, privacy set-up'),
          ('Months 3–5', 'Shadow pilot', 'Monsoon window: alerts reviewed by people before residents see them; resident drills'),
          ('Month 6', 'Review and scale', 'Report the KPIs, fix gaps, repeat in the next SDP township')]
x = 0.7
for i, (when, h, t) in enumerate(phases, 1):
    box(s, x, 2.7, 2.85, 1.7, fill=LIGHT)
    badge(s, x + 0.15, 2.8, i, d=0.34, size=12)
    tb(s, x + 0.6, 2.8, 2.15, 0.34, [{'runs': [(when, {'color': RED, 'bold': True})]}], size=12, anchor='m')
    tb(s, x + 0.15, 3.22, 2.55, 1.15, [{'text': h, 'bold': True, 'color': NAVY, 'size': 14, 'space_after': 2}, {'text': t, 'size': 11.5}])
    x += 3.017
blocks = [('Resources', ['Existing PWA and public data', 'Low-cost cloud; paid Gemini tier for real resident photos', 'Partners to request: SDP township team, local council, JPS, NADMA']),
          ('Privacy and safety', ['Consent, EXIF stripped, coarse location by default, retention limit', 'Decision-support only; official warnings stay with JPS and MetMalaysia']),
          ('Main risks and answers', ['Data access: fall back to public and synthetic data', 'Model error: show confidence, never a promise', 'False reports: threshold plus human review'])]
x = 0.7
for h, bl in blocks:
    box(s, x, 4.6, 3.85, 2.25, fill=WHITE, line='C9D1DA')
    tb(s, x + 0.18, 4.68, 3.5, 2.1, [{'text': h, 'bold': True, 'color': NAVY, 'size': 14, 'space_after': 4}] +
       [{'text': b, 'bullet': True, 'space_after': 3} for b in bl], size=11.5)
    x += 4.025

# =====================================================================
# 10. WHY WIN
# =====================================================================
s = content_slide('8 WHY OUR TEAM SHOULD WIN', 'Why our team should win',
                  source='Code repository and the 5-minute demo video are in the submission.',
                  notes='Close on honesty plus a ready pilot. 80 words, within the 100-word limit.')
box(s, 0.7, 1.5, 7.2, 5.1, fill=NAVY, radius=0.03)
tb(s, 1.0, 1.75, 6.6, 4.6, [
    'FloodWay turns a flood warning into a safe decision for one street. Our working PWA predicts risk, checks flood photos with AI before neighbours see them, and routes people to a relief centre.',
    {'text': 'We tested the market honestly: Waze, Google and MyDIMS cover parts of the chain, so we claim only the last mile and publish how we will measure it.', 'space_before': 10},
    {'text': 'Our data and assumptions are disclosed, and our pilot plan is ready for a township like Bandar Bukit Raja.', 'space_before': 10},
], size=20, color=WHITE, line=1.1)
tb(s, 8.3, 1.5, 4.3, 0.4, ['Disclosures'], size=18, bold=True, color=NAVY)
disc = [('Data', 'Public sources only, listed in Appendix A1. No SDP data used.'),
        ('AI use', 'Gemini Flash in the product; scikit-learn for the model; AI assistants used for research and drafting.'),
        ('IP', 'Existing FloodWay code is Background IP, licensed to SDP under the T&C.'),
        ('Links', 'Live prototype: floodways.netlify.app')]
y = 2.0
for h, t in disc:
    box(s, 8.3, y, 4.3, 1.08, fill=LIGHT)
    tb(s, 8.45, y + 0.07, 4.0, 1.0, [{'runs': [(h + '  ', {'bold': True, 'color': RED}), t]}], size=12, anchor='m')
    y += 1.18

# =====================================================================
# APPENDIX A1
# =====================================================================
s = content_slide('APPENDIX', 'Appendix A1: Data sources and assumptions',
                  source='SDP provided no datasets. Items marked "confirm" need the exact URL or licence from the team notebook before submission.',
                  notes='Organiser rule: all data sources, assumptions and simulated data must be stated.', appendix=True)
rows = [
    ('Item', 'Source', 'Licence or limit', 'Assumption or caveat'),
    ('Rainfall for model training', 'Malaysian rainfall 2000–2010, state and district, monthly (public file; confirm exact URL)', 'Confirm', 'Monthly resolution; period does not overlap the 2015–2021 labels'),
    ('Flood labels', 'JPS annual reports 2015–2021: locations, dates, geocoded points, extracted by the team', 'Public reports', 'Only reported incidents; under-reporting likely'),
    ('Live rainfall forecast', 'Real-time rainfall forecast API (integration ready; confirm provider)', 'Confirm', 'Not part of the accuracy figure'),
    ('Road network', 'OpenStreetMap', 'ODbL, attribution', 'Road completeness varies'),
    ('Relief centres (PPS)', 'NADMA and JKM public information, compiled by the team', 'No open dataset of locations found', 'Demo list; live open/closed status needs agency access'),
    ('Losses and deaths', 'DOSM special reports; police count (The Vibes)', 'Public statistics', '2022 figure omitted: DOSM revised it'),
    ('Survey', 'Team survey, 100 respondents, convenience sample', 'Own data', 'Not representative of ASEAN'),
    ('Simulation', 'Routing tested in synthetic flood scenarios built by the team', 'Own data', 'Not field evidence'),
    ('AI services', 'Gemini API (photo verification); scikit-learn (prediction)', 'Provider terms', 'Paid Gemini tier for real resident photos'),
]
colw = [2.2, 4.4, 2.1, 3.2]
ts = s.shapes.add_table(len(rows), 4, Inches(0.7), Inches(1.75), Inches(sum(colw)), Inches(4.9))
ts.name = 'Data register'
t = ts.table
t._tbl.tblPr.set('firstRow', '0')
t._tbl.tblPr.set('bandRow', '0')
for i, w in enumerate(colw):
    t.columns[i].width = Inches(w)
for r, row in enumerate(rows):
    t.rows[r].height = Inches(0.4 if r == 0 else 0.5)
    for c, v in enumerate(row):
        if r == 0:
            cell(t.cell(r, c), [v], NAVY, WHITE, 11, True)
        else:
            conf = 'confirm' in v.lower() or v == 'Confirm'
            cell(t.cell(r, c), [(v, {'bold': c == 0, 'color': NAVY if c == 0 else TEXT})], LIGHT if r % 2 else WHITE, TEXT, 10.5)

# =====================================================================
# APPENDIX A2
# =====================================================================
s = content_slide('APPENDIX', 'Appendix A2: Evaluation plan and known limitations',
                  source='Table cells marked TBD are filled from the team\'s own re-run; no figure here is estimated.',
                  notes='This slide shows we know where the model is weak and have a plan to fix it.', appendix=True)
tb(s, 0.7, 1.45, 6.2, 0.35, ['Planned evaluation (time-blocked split)'], size=15, bold=True, color=NAVY)
ev = [('Metric', 'Random Forest', 'No-flood baseline'),
      ('Accuracy', '74.55% (initial split)', 'TBD'),
      ('Precision / recall / F1', 'TBD', 'TBD'),
      ('PR-AUC', 'TBD', 'TBD'),
      ('Brier skill score', 'TBD', 'n/a'),
      ('Test-set flood share', 'TBD', 'n/a')]
ts = s.shapes.add_table(len(ev), 3, Inches(0.7), Inches(1.9), Inches(6.2), Inches(2.9))
ts.name = 'Evaluation table'
t = ts.table
t._tbl.tblPr.set('firstRow', '0')
t._tbl.tblPr.set('bandRow', '0')
for i, w in enumerate([2.4, 2.1, 1.7]):
    t.columns[i].width = Inches(w)
for r, row in enumerate(ev):
    t.rows[r].height = Inches(0.45)
    for c, v in enumerate(row):
        if r == 0:
            cell(t.cell(r, c), [v], NAVY, WHITE, 12, True)
        elif v == 'TBD':
            cell(t.cell(r, c), [(v, {'hl': YEL, 'bold': True, 'color': NAVY})], LIGHT if r % 2 else WHITE, NAVY, 12)
        else:
            cell(t.cell(r, c), [(v, {'bold': c == 0, 'color': NAVY if c == 0 else TEXT})], LIGHT if r % 2 else WHITE, TEXT, 12)
tb(s, 0.7, 5.0, 6.2, 1.8, [
    {'text': 'Method', 'bold': True, 'color': NAVY, 'size': 13, 'space_after': 3},
    {'text': 'Train on earlier years, test on later years; hold out whole districts', 'bullet': True, 'space_after': 2},
    {'text': 'Compare models with McNemar test and cluster bootstrap intervals', 'bullet': True, 'space_after': 2},
    {'text': 'Replace monthly rainfall with period-matched satellite or reanalysis rainfall (IMERG, ERA5)', 'bullet': True},
], size=12)
tb(s, 7.3, 1.45, 5.3, 0.35, ['Known limitations'], size=15, bold=True, color=NAVY)
lim = ['Rainfall file and flood labels cover different periods, and rainfall is monthly, so the model is a district-level susceptibility estimate until retrained',
       'Accuracy comes from one small initial split; Random Forest and Decision Tree are not clearly different',
       'Photo-verification accuracy is not yet measured; vision models can misjudge depth',
       'Routing was tested in simulation only; blocked-road behaviour in OSRM needs a field test',
       'Prototype was built on Gemini 1.5 Flash, which Google retired in Sep 2025; moving to the current Flash model is a pre-pilot task',
       'PPS open/closed status is not live; it needs agency data']
y = 1.9
for i, l in enumerate(lim, 1):
    box(s, 7.3, y, 5.3, 0.78, fill=LIGHT)
    badge(s, 7.4, y + 0.2, i, d=0.36, fill=NAVY, size=12)
    tb(s, 7.9, y, 4.6, 0.78, [l], size=11, anchor='m')
    y += 0.8

# =====================================================================
# APPENDIX A3
# =====================================================================
s = content_slide('APPENDIX', 'Appendix A3: Sources',
                  source='All pages accessed October 2026. Claims we could not source were removed from this deck.',
                  notes='Sources behind every number on the slides.', appendix=True)
srcs_l = [
    ('SDP University Hackathon page and T&C', 'https://simedarbyproperty.com/universityhackathon/'),
    ('SDP Sustainability Report 2023 (flood risk, wetlands)', 'https://www.simedarbyproperty.com/annualreport2023/download/sdp-sr2023.pdf'),
    ('Sinar Harian: Bukit Raja flash floods, May 2024', 'https://www.sinarharian.com.my/article/666626/edisi/selangor-kl/penduduk-bandar-bukit-raja-hadap-banjir-kilat-tiga-kali-setahun'),
    ('The Vibes: RM6.1bn loss, DOSM', 'https://www.thevibes.com/articles/business/53117/rm6.1-bil-loss-for-country-after-nationwide-floods-dosm'),
    ('The Vibes: 54 deaths, 5 Jan 2022', 'https://www.thevibes.com/articles/news/51355/floods-claim-54-lives-two-still-missing'),
    ('DOSM special report, floods 2024', 'https://www.dosm.gov.my/portal-main/release-content/special-report-on-impact-of-floods-in-malaysia-2024'),
    ('DOSM special report, floods 2025', 'https://www.dosm.gov.my/portal-main/release-content/special-report-on-impact-of-floods-in-malaysia2025'),
    ('WMO: 24 h warning cuts damage about 30%', 'https://public.wmo.int/node/25794'),
    ('WMO/UNDRR 2022: 8× higher mortality', 'https://wmo.int/node/15497'),
]
srcs_r = [
    ('Waze Help: flood hazard report', 'https://support.google.com/waze/answer/13809605'),
    ('Waze Help: road closures', 'https://support.google.com/waze/answer/13753511'),
    ('Google Flood Hub: covered locations', 'https://support.google.com/flood-hub/answer/16508958?hl=en'),
    ('Google Research: urban flash-flood model', 'https://research.google/blog/protecting-cities-with-ai-driven-flash-flood-forecasting/'),
    ('Malay Mail: NADMA MyDIMS launch', 'https://www.malaymail.com/news/malaysia/2026/04/24/nadma-launches-mydims-app-for-realtime-disaster-alerts-and-reporting/217562'),
    ('Malay Mail: JPS Kelantan forecast, Jan 2026', 'https://www.malaymail.com/news/malaysia/2026/01/04/heads-up-kelantan-did-warns-of-possible-floods-this-week-across-10-districts-from-jan-5-with-rising-rivers/204203'),
    ('NST: Waze routes men into floodwater, Nov 2024', 'https://master.nst.com.my/amp/news/nation/2024/11/1139451/three-rescued-after-waze-leads-them-floodwaters'),
    ('Gemini API changelog (1.5 Flash shutdown)', 'https://ai.google.dev/gemini-api/docs/changelog'),
    ('Open-Meteo historical API (ERA5)', 'https://open-meteo.com/en/docs/historical-weather-api'),
]
for xcol, lst in ((0.7, srcs_l), (6.85, srcs_r)):
    y = 1.5
    for lab, url in lst:
        tb(s, xcol, y, 5.75, 0.55, [{'text': lab, 'bold': True, 'color': NAVY, 'size': 11.5},
                                     {'runs': [(url if len(url) < 70 else url[:67] + '...', {'link': url, 'color': BLUE})], 'size': 9}])
        y += 0.58

prs.save(OUT)
print('saved', OUT, len(prs.slides._sldIdLst), 'slides')
