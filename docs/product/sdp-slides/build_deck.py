"""Rebuild the deck with slides 11 and 12 filled in.
Run from the repo root: python -I docs/product/sdp-slides/build_deck.py
Outputs (repo root): LarpGPT Slide-Sime Darby Slide Deck - updated.pdf / .pptx
The Canva pages stay as they are; only slides 11 and 12 change."""
import fitz
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor

SRC = "LarpGPT Slide-Sime Darby Slide Deck.pdf"
TPL = "SDP_University_Hackathon_TOWNSHIP-OF-TOMORROW-Template.pptx"
OUT = "LarpGPT Slide-Sime Darby Slide Deck - updated"
A = Path("docs/product/sdp-slides")
# page index -> (title png, body png, template slide index, title text)
PAGES = {
    10: ("title-11", "implementation", 14, "Four steps from sensor to street, gated before any resident is alerted"),
    11: ("title-12", "visual-concept", 15, "One app on every screen: predict, verify, evacuate"),
}
KICK = {10: "8 - IMPLEMENTATION", 11: "9 - VISUAL CONCEPT"}
RED, NAVY = RGBColor(0xC8, 0x10, 0x2E), RGBColor(0x14, 0x21, 0x3D)
BODY = fitz.Rect(75, 190, 1365, 730)  # deck content grid (1440 x 810 page)

doc = fitz.open(SRC)
for i, (title, body, _, _) in PAGES.items():
    pg = doc[i]
    pg.add_redact_annot(fitz.Rect(80, 80, 1300, 175))   # placeholder title
    pg.add_redact_annot(fitz.Rect(100, 255, 1300, 380))  # placeholder bullet
    pg.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE, graphics=fitz.PDF_REDACT_LINE_ART_NONE)
    pg.insert_image(pg.rect, filename=str(A / f"{title}.png"))
    pg.insert_image(BODY, filename=str(A / f"{body}.png"))
doc.save(OUT + ".pdf", garbage=3, deflate=True)

# PPTX: official template theme; Canva pages as full-bleed images, slides 11-12 native (editable title + picture)
prs = Presentation(TPL)
sections = prs.slide_layouts.get_by_name("Section Header")
tpl_slides = list(prs.slides)
ids = prs.slides._sldIdLst
for sid in list(ids):
    prs.part.drop_rel(sid.rId); ids.remove(sid)
W, H = prs.slide_width, prs.slide_height
tmp = Path(OUT + "_img"); tmp.mkdir(exist_ok=True)
for i, page in enumerate(fitz.open(SRC)):
    s = prs.slides.add_slide(sections)
    if i in PAGES:
        _, body, _, text = PAGES[i]
        for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
        for txt, (l, t, w, h), pt, bold, rgb in ((KICK[i], (59, 52, 500, 28), 12, True, RED), (text, (82, 80, 1130, 90), 23, True, NAVY)):
            tb = s.shapes.add_textbox(Emu(int(W * l / 1440)), Emu(int(H * t / 810)), Emu(int(W * w / 1440)), Emu(int(H * h / 810)))
            tb.text_frame.word_wrap = True
            r = tb.text_frame.paragraphs[0].add_run(); r.text = txt
            r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(pt), bold, "Outfit", rgb
        x, y = BODY.x0, BODY.y0
        s.shapes.add_picture(str(A / f"{body}.png"), Emu(int(W * x / 1440)), Emu(int(H * y / 810)), width=Emu(int(W * BODY.width / 1440)))
    else:
        for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
        png = tmp / f"p{i+1:02d}.png"
        page.get_pixmap(dpi=144).save(png)
        s.shapes.add_picture(str(png), 0, 0, W, H)
prs.save(OUT + ".pptx")
for f in tmp.iterdir(): f.unlink()
tmp.rmdir()
print("ok")
